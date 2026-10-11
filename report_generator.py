"""Генерация отчётов в .docx и .xlsx."""
from docx import Document
from openpyxl import Workbook


def save_to_docx(data: dict, filepath: str) -> None:
    """Сохранить результаты расчёта в .docx."""
    doc = Document()
    doc.add_heading("Отчёт по банковскому расчёту", level=1)

    doc.add_heading("Параметры", level=2)
    for key, value in data["params"].items():
        doc.add_paragraph(f"{key}: {value}")

    doc.add_heading("Результаты", level=2)
    for key, value in data["results"].items():
        doc.add_paragraph(f"{key}: {value}")

    if data.get("schedule"):
        doc.add_heading("График платежей", level=2)
        table = doc.add_table(rows=1, cols=len(data["schedule"][0]))
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, col in enumerate(data["schedule"][0].keys()):
            hdr[i].text = str(col)
        for row in data["schedule"]:
            cells = table.add_row().cells
            for i, value in enumerate(row.values()):
                cells[i].text = str(value)

    doc.save(filepath)


def save_to_xlsx(data: dict, filepath: str) -> None:
    """Сохранить результаты расчёта в .xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    ws.append(["Параметр", "Значение"])
    for key, value in data["params"].items():
        ws.append([key, value])
    for key, value in data["results"].items():
        ws.append([key, value])

    if data.get("schedule"):
        ws2 = wb.create_sheet("График")
        headers = list(data["schedule"][0].keys())
        ws2.append(headers)
        for row in data["schedule"]:
            ws2.append(list(row.values()))

    wb.save(filepath)